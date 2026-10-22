#!/usr/bin/env python3
"""
Build the GSA Products catalog data from your price list.

Usage:
    python tools/build_catalog.py catalog/pricelist.xlsx
    python tools/build_catalog.py catalog/pricelist.csv
    python tools/build_catalog.py catalog/pricelist.zip   (zip holding the .xlsx or .csv)

Reads a spreadsheet with these columns (header names, any order):
    MAS Contract, manufacturer, manufacturer_part_number, vendor_part_number,
    sin, item_name, item_description, uom, Price, country_of_origin, delivery_days

Writes small JSON files into  data/  that products.html loads on demand:
    data/meta.json        counts, manufacturer list, SIN list
    data/c/<n>.json       items, 500 per file
    data/i/<key>.json     search index (by first 3 letters of each word)
    data/m/<n>.json       item ids per manufacturer
    data/s/<sin>.json     item ids per SIN
    data/p/<key>.json     part number -> item id (for product links)
"""
import csv, json, os, re, shutil, sys, time
from collections import defaultdict

CHUNK = 500
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data")

# header in spreadsheet  ->  short key used on the site
COLS = {
    "mascontract": "contract",
    "manufacturer": "mfr",
    "manufacturerpartnumber": "mpn",
    "vendorpartnumber": "vpn",
    "sin": "sin",
    "itemname": "name",
    "itemdescription": "desc",
    "uom": "uom",
    "price": "price",
    "govtpricewithfee": "price",
    "defaultphoto": "photo",
    "countryoforigin": "coo",
    "deliverydays": "days",
}
FIELDS = ["contract", "mfr", "mpn", "vpn", "sin", "name", "desc", "uom", "price", "coo", "days", "photo"]


def norm_header(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def fmt_price(v):
    s = clean(v).replace("$", "").replace(",", "")
    try:
        return f"{float(s):,.2f}"
    except ValueError:
        return clean(v)


def read_rows(path):
    if path.lower().endswith(".zip"):
        import zipfile, tempfile
        z = zipfile.ZipFile(path)
        name = next(n for n in z.namelist() if n.lower().endswith((".xlsx", ".xlsm", ".csv")) and not n.startswith("__MACOSX"))
        tmp = os.path.join(tempfile.mkdtemp(), os.path.basename(name))
        with open(tmp, "wb") as f:
            f.write(z.read(name))
        yield from read_rows(tmp)
        return
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        for row in it:
            yield header, row
    else:
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            r = csv.reader(f)
            header = next(r)
            for row in r:
                yield header, row


def words(text):
    return [w for w in re.split(r"[^a-z0-9]+", text.lower()) if len(w) >= 2]


def pn_key(pn):
    return re.sub(r"[^a-z0-9]", "", pn.lower())


def shard(token):
    return token[:3]


def main(src):
    t0 = time.time()
    items, colmap = [], None
    pdir = os.path.join(os.path.dirname(OUT), "img", "p")
    photos = set(os.listdir(pdir)) if os.path.isdir(pdir) else set()
    photo_missing = 0
    for header, row in read_rows(src):
        if colmap is None:
            colmap = {}
            for i, h in enumerate(header):
                k = COLS.get(norm_header(h))
                if k:
                    colmap[k] = i
            missing = [c for c in ("name", "mpn", "price") if c not in colmap]
            if missing:
                sys.exit(f"Missing required column(s): {missing}. Found: {list(header)}")
        rec = {k: clean(row[i]) if i < len(row) else "" for k, i in colmap.items()}
        if not (rec.get("name") or rec.get("mpn")):
            continue  # skip blank lines
        rec["price"] = fmt_price(rec.get("price", ""))
        ph = rec.get("photo", "")
        if ph:
            webp = ph.rsplit(".", 1)[0].lower() + ".webp"
            if webp in photos:
                rec["photo"] = webp
            else:
                rec["photo"] = ""
                photo_missing += 1
        items.append([rec.get(f, "") for f in FIELDS])

    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    for d in ("c", "i", "m", "s", "p"):
        os.makedirs(os.path.join(OUT, d))

    index = defaultdict(lambda: defaultdict(list))  # shard -> token -> ids
    by_mfr, by_sin = defaultdict(list), defaultdict(list)
    pn = defaultdict(dict)

    F = {f: n for n, f in enumerate(FIELDS)}
    for i, it in enumerate(items):
        # main words: name, manufacturer, part numbers. Description-only words are
        # stored with a "~" prefix so search can rank name matches higher.
        toks = set(words(" ".join([it[F["name"]], it[F["mfr"]], it[F["mpn"]], it[F["vpn"]]])))
        for p in (it[F["mpn"]], it[F["vpn"]]):
            k = pn_key(p)
            if len(k) >= 2:
                toks.add(k)
                lst = pn[k[:2]].setdefault(k, [])
                if i not in lst:
                    lst.append(i)
        for t in toks:
            index[shard(t)][t].append(i)
        for t in set(words(it[F["desc"]])) - toks:
            index[shard(t)]["~" + t].append(i)
        if it[F["mfr"]]:
            by_mfr[it[F["mfr"]]].append(i)
        if it[F["sin"]]:
            by_sin[it[F["sin"]]].append(i)

    def dump(path, obj):
        with open(os.path.join(OUT, path), "w", encoding="utf-8") as f:
            json.dump(obj, f, separators=(",", ":"), ensure_ascii=False)

    for n in range(0, len(items), CHUNK):
        dump(f"c/{n // CHUNK}.json", items[n:n + CHUNK])
    for k, toks in index.items():
        dump(f"i/{k}.json", toks)
    mfrs = sorted(by_mfr, key=str.lower)
    for n, m in enumerate(mfrs):
        dump(f"m/{n}.json", by_mfr[m])
    sins = sorted(by_sin)
    for s in sins:
        dump(f"s/{re.sub(r'[^A-Za-z0-9_-]', '_', s)}.json", by_sin[s])
    for k, d in pn.items():
        dump(f"p/{k}.json", {pn_: (ids[0] if len(ids) == 1 else ids) for pn_, ids in d.items()})

    featured = []
    fpath = os.path.join(os.path.dirname(OUT), "catalog", "featured.txt")
    if os.path.isfile(fpath):
        by_mpn = {}
        for n, it in enumerate(items):
            by_mpn.setdefault(it[F["mpn"]].strip().lower(), n)
        for line in open(fpath, encoding="utf-8"):
            pn_ = line.strip()
            if not pn_ or pn_.startswith("#"):
                continue
            n = by_mpn.get(pn_.lower())
            if n is None:
                print(f"WARNING: featured part number not found: {pn_}")
            else:
                featured.append(n)

    dump("meta.json", {
        "featured": featured,
        "count": len(items),
        "chunk": CHUNK,
        "fields": FIELDS,
        "manufacturers": [[m, len(by_mfr[m])] for m in mfrs],
        "sins": [[s, len(by_sin[s])] for s in sins],
        "built": time.strftime("%Y-%m-%d"),
    })

    warn = [(n + 2, it[F["mpn"]], it[F["days"]]) for n, it in enumerate(items)
            if it[F["days"]] and (not it[F["days"]].isdigit() or int(it[F["days"]]) > 365)]
    if warn:
        print(f"WARNING: {len(warn)} row(s) have an unusual delivery_days value (shown as-is on the site):")
        for row, mpn, d in warn[:20]:
            print(f"  spreadsheet row {row}: {mpn}  delivery_days={d}")

    files = sum(len(fs) for _, _, fs in os.walk(OUT))
    size = sum(os.path.getsize(os.path.join(r, f)) for r, _, fs in os.walk(OUT) for f in fs)
    print(f"{len(items):,} items, {len(mfrs):,} manufacturers, {len(sins)} SINs")
    with_photo = sum(1 for it in items if it[F["photo"]])
    print(f"{with_photo:,} items with a photo; {photo_missing:,} rows name a photo that is not in img/p")
    print(f"{files:,} files, {size / 1e6:.1f} MB, {time.time() - t0:.1f}s")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
